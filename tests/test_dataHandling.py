import csv
import json
import os

from dotenv import load_dotenv
from openpyxl import load_workbook
import pytest

from utils.jsonHandling import jsonHandling


def test_jsonHandling():
    with open('testData\\creds.json') as data:
        creds = json.load(data)
        print(creds['email'])

def test_2():
    data = jsonHandling('testData\\creds.json')


def test_csvHandling():
     with open('testData\\credentails.csv') as data:
            creds = csv.DictReader(data)
            values =[]
            for i in creds:
                 values.append(i)
            print(values[1]['username'])



def test_handlinhExcel():
     workbook = load_workbook("testData\\sample_creds.xlsx")
     sheet = workbook["Sheet2"]
     values =[]
     for i in sheet.iter_rows(min_row=2,values_only=True):
          values.append(i)

# @pytest.mark.dh
def test_cLI():
     username = os.getenv("usName_k")
     pasword =  os.getenv("pw_k")
     print(username)
     print(pasword)


@pytest.mark.dh
def test_env():
     load_dotenv(os.getenv("file"))
     # load_dotenv(".env")
     username = os.getenv("key1")
     pasword =  os.getenv("key2")
     print(username)
     print(pasword)

     

